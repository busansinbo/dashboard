"""Extract official 부산 small-business quarterly totals from Big-데이터웨이브 workbooks."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

import openpyxl

from external_indicators_common import ROOT, append_log, load_config, load_data, merge_period_rows, numeric, period_quarter, utc_now, write_json


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def decimal_places(number_format: str) -> int:
    match = re.search(r"0\.(0+)", number_format or "")
    return len(match.group(1)) if match else 0


def find_summary_sheets(workbook: openpyxl.Workbook) -> tuple[str, str] | None:
    basic = next((name for name in workbook.sheetnames if name.startswith("1.") and "사업체 기본현황" in name), None)
    opening = next((name for name in workbook.sheetnames if name.startswith("5.") and "창폐업" in name), None)
    return (basic, opening) if basic and opening else None


def quarter_or_none(value: Any) -> str | None:
    text = str(value or "")
    if "분기" not in text or "연간" in text:
        return None
    try:
        return period_quarter(text)
    except ValueError:
        return None


def is_period_group(value: Any) -> bool:
    """Return True for both quarterly and annual column-group headers."""
    text = str(value or "").replace(" ", "")
    return "분기" in text or "연간" in text


def parse_summary_workbook(path: Path) -> dict[str, dict[str, Any]] | None:
    preview = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pair = find_summary_sheets(preview)
    if not pair:
        preview.close(); return None
    preview.close()
    values = openpyxl.load_workbook(path, read_only=False, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=False, data_only=False)
    basic_name, opening_name = pair
    basic, opening, opening_formula = values[basic_name], values[opening_name], formulas[opening_name]
    digest, retrieved = file_hash(path), utc_now()

    small_col = next(c for c in range(1, basic.max_column + 1) if str(basic.cell(3, c).value or "").strip() == "소상공인")
    next_group = next((c for c in range(small_col + 1, basic.max_column + 2) if c > basic.max_column or basic.cell(3, c).value is not None), basic.max_column + 1)
    business_col = next(c for c in range(small_col, next_group) if str(basic.cell(4, c).value or "").replace(" ", "") == "사업체수")
    sales_col = next(c for c in range(small_col, next_group) if str(basic.cell(4, c).value or "").replace(" ", "") == "매출액")
    rows: dict[str, dict[str, Any]] = {}
    for row_number in range(5, basic.max_row + 1):
        raw_period = basic.cell(row_number, 1).value
        period = quarter_or_none(raw_period)
        if not period:
            continue
        rows[period] = {
            "period": period, "periodEnd": f"{period[:4]}-{int(period[-1]) * 3:02d}",
            "businessCount": numeric(basic.cell(row_number, business_col).value), "salesAmount": numeric(basic.cell(row_number, sales_col).value),
            "status": "provisional" if "P" in str(raw_period).upper() else "final",
            "units": {"businessCount": "개", "salesAmount": "백만원", "startupRate": "%", "closureRate": "%"},
            "region": "부산 전체", "aggregationBasis": "official_busan_total", "source": "부산 Big-데이터웨이브",
            "sourceFile": path.name, "sourceSheet": basic_name, "sourceRows": {"businessCount": row_number, "salesAmount": row_number},
            "fileSha256": digest, "retrievedAt": retrieved, "sourceType": {"businessCount": "published", "salesAmount": "published"},
        }

    total_row = next(r for r in range(1, opening.max_row + 1) if str(opening.cell(r, 1).value or "").replace(" ", "") == "소상공인")
    # Annual totals can immediately follow Q4 and repeat the same labels.  The
    # annual group must therefore delimit the Q4 group even though it is not a
    # quarterly observation.  Otherwise its cumulative values overwrite the
    # true Q4 columns in labels4.
    group_starts = [c for c in range(1, opening.max_column + 1) if is_period_group(opening.cell(3, c).value)]
    headers = [(c, quarter_or_none(opening.cell(3, c).value)) for c in group_starts]
    headers = [(c, period) for c, period in headers if period]
    for start_col, period in headers:
        if period not in rows:
            continue
        end_col = next((c for c in group_starts if c > start_col), opening.max_column + 1)
        labels4 = {str(opening.cell(4, c).value or "").replace(" ", ""): c for c in range(start_col, end_col)}
        activity_col = labels4.get("활동사업체수") or labels4.get("사업체수") or labels4.get("활동")
        startup_col, closure_col = labels4.get("창업수"), labels4.get("폐업수")
        if not all((activity_col, startup_col, closure_col)):
            raise ValueError(f"{path.name} {period}: activity/startup/closure columns are incomplete")
        activity, startup, closure = (numeric(opening.cell(total_row, c).value) for c in (activity_col, startup_col, closure_col))
        calculated = {"startupRate": startup / activity * 100, "closureRate": closure / activity * 100}
        rate_columns: dict[str, int] = {}
        for c in range(start_col, end_col):
            label = str(opening.cell(5, c).value or "").replace(" ", "")
            if label == "창업률": rate_columns["startupRate"] = c
            if label == "폐업률": rate_columns["closureRate"] = c
        reference_row = next((r for r in range(total_row + 1, opening.max_row + 1) if "창업률" in str(opening.cell(r, 1).value or "") and "폐업률" in str(opening.cell(r, 1).value or "")), None)
        if reference_row:
            if opening_formula.cell(reference_row, startup_col).data_type == "f": rate_columns.setdefault("startupRate", startup_col)
            if opening_formula.cell(reference_row, closure_col).data_type == "f": rate_columns.setdefault("closureRate", closure_col)

        row = rows[period]
        row.update({"activityBusinessCount": activity, "startupCount": startup, "closureCount": closure})
        row["sourceRows"].update({"activityBusinessCount": total_row, "startupCount": total_row, "closureCount": total_row})
        row["sourceType"].update({"activityBusinessCount": "published", "startupCount": "published", "closureCount": "published"})
        row["formula"] = {"startupRate": "startupCount / activityBusinessCount * 100", "closureRate": "closureCount / activityBusinessCount * 100"}
        row["formulaInputs"] = {"activityBusinessCount": activity, "startupCount": startup, "closureCount": closure}
        row["publishedRateValidation"] = {}
        for key, label in (("startupRate", "창업률"), ("closureRate", "폐업률")):
            published_col = rate_columns.get(key)
            if published_col and str(opening.cell(5, published_col).value or "").replace(" ", "") == label:
                published_row = total_row
            else:
                published_row = reference_row
            published = numeric(opening.cell(published_row, published_col).value) if published_col and published_row else None
            if published is not None:
                places = decimal_places(opening_formula.cell(published_row, published_col).number_format)
                if abs(published - calculated[key]) > 1e-10 or round(published, places) != round(calculated[key], places):
                    raise ValueError(f"{path.name} {period}: published {key} does not reconcile")
                row[key], row["sourceType"][key] = published, "published"
                row["publishedRateValidation"][key] = {"published": published, "recalculated": calculated[key], "difference": published - calculated[key], "displayDecimals": places, "displayMatched": True}
            else:
                row[key], row["sourceType"][key] = calculated[key], "calculated_from_official_formula"
    values.close(); formulas.close()
    return rows


def consolidate(parsed: list[tuple[Path, dict[str, dict[str, Any]]]]) -> tuple[list[dict[str, Any]], list[str]]:
    parsed.sort(key=lambda item: max(item[1]), reverse=True)
    result: dict[str, dict[str, Any]] = {}; supporting: dict[str, list[str]] = {}
    for path, rows in parsed:
        for period, row in rows.items():
            supporting.setdefault(period, []).append(path.name)
            if period not in result:
                result[period] = row
            else:
                for key in ("startupRate", "closureRate"):
                    if (row["sourceType"].get(key) == "published" and row.get(key) is not None
                            and result[period].get(key) is not None and abs(row[key] - result[period][key]) <= 1e-10):
                        result[period][key] = row[key]; result[period]["sourceType"][key] = "published"
                        result[period]["publishedRateValidation"][key] = row["publishedRateValidation"][key]
                        result[period].setdefault("rateSourceFiles", {})[key] = path.name
    for period, row in result.items():
        row["supportingSourceFiles"] = supporting[period]
        row["sourceValues"] = {key: row[key] for key in ("businessCount", "salesAmount", "startupRate", "closureRate")}
    return [result[key] for key in sorted(result)], [path.name for path, _ in parsed]


def main() -> int:
    parser = argparse.ArgumentParser(); parser.add_argument("files", nargs="*"); parser.add_argument("--dry-run", action="store_true"); args = parser.parse_args()
    config = load_config(); qconfig = config["quarterly"]["busanSmallBusiness"]; source_dir = ROOT / qconfig["sourceDirectory"]
    paths = [Path(value).resolve() for value in args.files] if args.files else sorted(path for path in source_dir.glob("*.xlsx") if not path.name.startswith("~$"))
    parsed, inspected, unsupported = [], [], []
    for path in paths:
        inspected.append(path.name); rows = parse_summary_workbook(path)
        if rows: parsed.append((path, rows))
        else: unsupported.append(path.name)
    if not parsed:
        print("No supported summary workbooks found"); return 0
    incoming, source_files = consolidate(parsed); data = load_data(config)
    merged, changes = merge_period_rows(data["quarterly"].get("busanSmallBusiness", []), incoming)
    data["quarterly"]["busanSmallBusiness"] = merged; data["generatedAt"] = utc_now(); data["latest"]["quarterlyPeriod"] = max(row["period"] for row in merged)
    data["sources"]["busanSmallBusiness"].update({"retrievedAt": data["generatedAt"], "sourceFile": source_files[0], "sourceFiles": source_files,
        "inspectedFiles": inspected, "formula": {"startupRate": "창업수/활동사업체수*100", "closureRate": "폐업수/활동사업체수*100"}})
    if not args.dry_run:
        write_json(ROOT / config["output"], data)
        append_log(config, {"type": "quarterly-excel", "status": "success", "files": inspected, "unsupportedDetailedFiles": unsupported,
            "periodRange": [incoming[0]["period"], incoming[-1]["period"]],
            "reason": "Q4 quarterly startup/closure values corrected by separating adjacent annual cumulative column groups",
            "changes": changes})
    print(json.dumps({"status": "ok", "dryRun": args.dry_run, "inspected": len(inspected), "supportedSummaryFiles": len(parsed), "periods": len(incoming),
        "first": incoming[0]["period"], "last": incoming[-1]["period"], "changes": len(changes)}, ensure_ascii=False)); return 0


if __name__ == "__main__":
    raise SystemExit(main())
