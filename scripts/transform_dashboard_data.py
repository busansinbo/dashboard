#!/usr/bin/env python3
"""Convert the immutable source ZIP into standardized dashboard JSON.

The archive is read only. Outputs are written to a separate directory.
"""

from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import json
import re
import tempfile
import warnings
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DIMENSIONS = {
    "district": "district",
    "age": "age_group",
    "industry": "industry_major",
    "tenure": "business_age",
}
METRICS = (
    "new_guarantee_count",
    "new_guarantee_amount",
    "accident_net_count",
    "accident_net_amount",
    "guarantee_balance_count",
    "guarantee_balance_amount",
)
FILL_COLUMNS = (
    "period", "province", "district", "industry_major", "industry_middle",
    "industry_detail", "age_group", "business_age",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def decimal_value(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value).replace(",", "").strip())


def json_number(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def period_from_filename(path: Path) -> str | None:
    patterns = (
        r"_(\d{2})(\d{2})\.xlsx$",
        r"_(20\d{2})(\d{2})\.xlsx$",
        r"_(20\d{2})-(\d{2})\.xlsx$",
    )
    for pattern in patterns:
        match = re.search(pattern, path.name)
        if match:
            year = match.group(1)
            year = f"20{int(year):02d}" if len(year) == 2 else year
            return f"{year}-{int(match.group(2)):02d}"
    return None


def normalize_period(value: Any) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) != 6:
        raise ValueError(f"올바르지 않은 기준년월: {value!r}")
    return f"{digits[:4]}-{digits[4:]}"


def clean_text(value: Any) -> str:
    return str(value).strip() if value not in (None, "") else ""


@dataclass
class Totals:
    values: dict[str, Decimal] = field(default_factory=lambda: {key: Decimal("0") for key in METRICS})

    def add(self, record: dict[str, Any]) -> None:
        for key in METRICS:
            self.values[key] += record.get(key, Decimal("0"))

    def as_json(self) -> dict[str, int | float]:
        return {key: json_number(value) for key, value in self.values.items()}


class Aggregator:
    def __init__(
        self,
        age_mapping: dict[str, str] | None = None,
        target_province: str = "부산",
        target_districts: Iterable[str] | None = None,
    ) -> None:
        self.age_mapping = age_mapping or {}
        self.target_province = target_province
        self.target_districts = set(target_districts or ())
        self.monthly: dict[str, Totals] = defaultdict(Totals)
        self.scope_monthly: dict[str, dict[str, Totals]] = {
            "busan_16_districts": defaultdict(Totals),
            "outside_busan": defaultdict(Totals),
            "busan_other": defaultdict(Totals),
        }
        self.dimension: dict[str, dict[tuple[str, str], Totals]] = {
            name: defaultdict(Totals) for name in DIMENSIONS
        }

    def add(self, record: dict[str, Any]) -> None:
        period = record["period"]
        self.monthly[period].add(record)
        is_busan_district = (
            record.get("province") == self.target_province
            and record.get("district") in self.target_districts
        )
        if is_busan_district:
            self.scope_monthly["busan_16_districts"][period].add(record)
        elif record.get("province") != self.target_province:
            self.scope_monthly["outside_busan"][period].add(record)
        else:
            self.scope_monthly["busan_other"][period].add(record)
        for output_name, field_name in DIMENSIONS.items():
            if output_name == "district" and not is_busan_district:
                continue
            label = record.get(field_name) or "미분류"
            if output_name == "age":
                label = self.age_mapping.get(label, label)
            self.dimension[output_name][(period, label)].add(record)


def find_header(ws, required: list[str]) -> tuple[int, list[str]]:
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        headers = [clean_text(value) for value in row]
        if set(required).issubset(set(headers)):
            return row_number, headers
    raise ValueError(f"필수 열을 포함하는 헤더 행을 찾지 못했습니다: {required}")


def records_from_workbook(
    path: Path,
    dataset_type: str,
    required_headers: list[str],
    header_map: dict[str, str],
    fixed_period: str | None,
) -> Iterable[tuple[int, dict[str, Any]]]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"예상과 다른 시트 수({len(workbook.sheetnames)}): {workbook.sheetnames}")
        ws = workbook[workbook.sheetnames[0]]
        header_row, raw_headers = find_header(ws, required_headers)
        indexes = {header_map[name]: index for index, name in enumerate(raw_headers) if name in header_map}
        carried = {key: "" for key in FILL_COLUMNS}
        for source_row, values in enumerate(
            ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            if not any(value not in (None, "") for value in values):
                continue
            record: dict[str, Any] = {"dataset_type": dataset_type}
            for standard_name, index in indexes.items():
                value = values[index] if index < len(values) else None
                if standard_name in FILL_COLUMNS:
                    text = clean_text(value)
                    if text:
                        carried[standard_name] = text
                    record[standard_name] = carried[standard_name]
                elif standard_name in METRICS:
                    record[standard_name] = decimal_value(value)
            record["period"] = fixed_period or normalize_period(record.get("period"))
            for metric in METRICS:
                record.setdefault(metric, Decimal("0"))
            yield source_row, record
    finally:
        workbook.close()


def choose_source_files(root: Path) -> tuple[list[tuple[Path, str]], list[tuple[Path, str | None]], list[str]]:
    warnings_out: list[str] = []
    balance_by_period: dict[str, list[Path]] = defaultdict(list)
    for path in (root / "잔액").glob("*.xlsx"):
        period = period_from_filename(path)
        if period:
            balance_by_period[period].append(path)
    balance_files: list[tuple[Path, str]] = []
    for period, paths in sorted(balance_by_period.items()):
        if len(paths) > 1:
            raise ValueError(f"잔액 파일 중복({period}): {[p.name for p in paths]}")
        balance_files.append((paths[0], period))
    accident_files = sorted((root / "사고").glob("*.xlsx"), key=lambda p: p.stat().st_mtime)
    if not accident_files:
        raise ValueError("사고 Excel 파일이 없습니다.")
    monthly_accident: dict[str, list[Path]] = defaultdict(list)
    integrated_accident: list[Path] = []
    for path in accident_files:
        period = period_from_filename(path)
        if period:
            monthly_accident[period].append(path)
        else:
            integrated_accident.append(path)
    selected_accident: list[tuple[Path, str | None]] = []
    if monthly_accident:
        for period, paths in sorted(monthly_accident.items()):
            if len(paths) > 1:
                raise ValueError(f"사고 파일 중복({period}): {[p.name for p in paths]}")
            selected_accident.append((paths[0], period))
        if integrated_accident:
            warnings_out.append("월별 사고 파일을 사용하고 기존 통합 사고 파일은 제외했습니다.")
    else:
        if len(integrated_accident) > 1:
            ignored = [p.name for p in integrated_accident[:-1]]
            warnings_out.append(f"통합 사고 파일이 여러 개여서 최신 파일만 사용: 제외={ignored}")
        selected_accident.append((integrated_accident[-1], None))
    return balance_files, selected_accident, warnings_out


def annualized_accident_rate(period: str, current: Totals, history: Iterable[tuple[str, Totals]]) -> float | None:
    year, month = map(int, period.split("-"))
    balance = current.values["guarantee_balance_amount"]
    if not balance:
        return None
    ytd_net = sum(
        (totals.values["accident_net_amount"] for item_period, totals in history if item_period[:4] == period[:4] and item_period <= period),
        Decimal("0"),
    )
    annualization_days = Decimal("365")
    month_end = date(year, month, calendar.monthrange(year, month)[1])
    elapsed_days = Decimal(str(month_end.timetuple().tm_yday))
    return float(ytd_net / balance * (annualization_days / elapsed_days) * 100)


def dashboard_row(label: str, totals: Totals, rate: float | None, period: str | None = None) -> dict[str, Any]:
    value = totals.values
    row = {
        "label": label,
        "newGuarantee": json_number(value["new_guarantee_amount"]),
        "accidents": json_number(value["accident_net_amount"]),
        "accidentRate": round(rate, 4) if rate is not None else None,
        "balance": json_number(value["guarantee_balance_amount"]),
    }
    if period:
        row["period"] = period
    return row


def build_outputs(
    aggregator: Aggregator,
    config: dict[str, Any],
    source_hash: str,
    file_reports: list[dict[str, Any]],
    warnings_out: list[str],
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    periods = sorted(aggregator.monthly)
    if not periods:
        raise ValueError("변환 대상 데이터가 없습니다.")
    latest = periods[-1]
    monthly_history = list(aggregator.monthly.items())
    latest_year, latest_month = map(int, latest.split("-"))
    dashboard = {
        "meta": {
            "title": "부산 신용보증 통계",
            "year": int(latest[:4]),
            "updatedAt": date(latest_year, latest_month, calendar.monthrange(latest_year, latest_month)[1]).isoformat(),
            "periodLabel": f"{latest[:4]}년 {int(latest[5:])}월 기준",
            "unit": config.get("amount_unit") or "원본 금액 단위(확인 필요)",
            "sourceSha256": source_hash,
            "dataType": "actual",
            "aggregationScope": {
                "overall": "부산신용보증재단 관리 전체 보증(지역 제한 없음)",
                "district": "부산광역시 16개 구·군",
                "ageIndustryTenure": "부산신용보증재단 관리 전체 보증(지역 제한 없음)",
            },
            "excludedIndicators": [
                {
                    "name": "보증공급",
                    "reason": "현재 원본에 독립된 보증공급 컬럼과 공식 산출 기준이 없어 신규보증 실행으로 대체하지 않음"
                }
            ],
        },
        "monthly": [
            dashboard_row(
                f"{int(period[5:])}월",
                aggregator.monthly[period],
                annualized_accident_rate(period, aggregator.monthly[period], monthly_history),
                period,
            )
            for period in periods
        ],
    }
    for output_name in DIMENSIONS:
        rows = [
            dashboard_row(
                label,
                totals,
                annualized_accident_rate(
                    latest,
                    totals,
                    [(item_period, item_totals) for (item_period, item_label), item_totals in aggregator.dimension[output_name].items() if item_label == label],
                ),
            )
            for (period, label), totals in aggregator.dimension[output_name].items()
            if period == latest
        ]
        if output_name == "age":
            age_order = {label: index for index, label in enumerate(config.get("age_group_order", []))}
            dashboard[output_name] = sorted(rows, key=lambda row: age_order.get(row["label"], 999))
        else:
            dashboard[output_name] = sorted(rows, key=lambda row: row["newGuarantee"], reverse=True)

    history = {
        "meta": dashboard["meta"],
        "monthly": [
            {"period": period, **aggregator.monthly[period].as_json()} for period in periods
        ],
        "dimensions": {
            output_name: [
                {"period": period, "label": label, **totals.as_json()}
                for (period, label), totals in sorted(values.items())
            ]
            for output_name, values in aggregator.dimension.items()
        },
    }

    validation_files = []
    overall_source = Totals()
    overall_converted = Totals()
    source_rows = converted_rows = 0
    for item in file_reports:
        source_rows += item["selected_rows"]
        converted_rows += item["converted_rows"]
        for key in METRICS:
            overall_source.values[key] += Decimal(str(item["source_totals"][key]))
            overall_converted.values[key] += Decimal(str(item["converted_totals"][key]))
        item["row_count_match"] = item["selected_rows"] == item["converted_rows"]
        item["sum_match"] = item["source_totals"] == item["converted_totals"]
        validation_files.append(item)
    validation = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_sha256": source_hash,
        "source_unchanged": True,
        "target_province": config["target_province"],
        "target_districts": config["target_districts"],
        "aggregation_scope": dashboard["meta"]["aggregationScope"],
        "warnings": warnings_out,
        "files": validation_files,
        "overall": {
            "source_selected_rows": source_rows,
            "converted_rows": converted_rows,
            "row_count_match": source_rows == converted_rows,
            "source_totals": overall_source.as_json(),
            "converted_totals": overall_converted.as_json(),
            "sum_match": overall_source.as_json() == overall_converted.as_json(),
        },
        "regional_scope": {
            scope: {
                period: totals.as_json() for period, totals in sorted(periods_map.items())
            }
            for scope, periods_map in aggregator.scope_monthly.items()
        },
    }
    latest_dashboard_row = next(row for row in dashboard["monthly"] if row["period"] == latest)
    latest_totals = aggregator.monthly[latest]
    metric_display_map = {
        "new_guarantee_amount": "newGuarantee",
        "accident_net_amount": "accidents",
        "guarantee_balance_amount": "balance",
    }
    kpi_checks = []
    for source_metric, display_metric in metric_display_map.items():
        source_value = json_number(latest_totals.values[source_metric])
        display_value = latest_dashboard_row[display_metric]
        kpi_checks.append({
            "indicator": display_metric,
            "source_aggregate": source_value,
            "dashboard_value": display_value,
            "match": source_value == display_value,
        })
    expected_rate = annualized_accident_rate(latest, latest_totals, monthly_history)
    displayed_rate = latest_dashboard_row["accidentRate"]
    kpi_checks.append({
        "indicator": "accidentRate",
        "source_aggregate": round(expected_rate, 4) if expected_rate is not None else None,
        "dashboard_value": displayed_rate,
        "match": (round(expected_rate, 4) if expected_rate is not None else None) == displayed_rate,
    })
    dimension_checks = []
    for dimension_name in DIMENSIONS:
        rows = dashboard[dimension_name]
        for source_metric, display_metric in metric_display_map.items():
            expected_totals = (
                aggregator.scope_monthly["busan_16_districts"][latest]
                if dimension_name == "district"
                else latest_totals
            )
            source_value = json_number(expected_totals.values[source_metric])
            display_total = sum(Decimal(str(row[display_metric])) for row in rows)
            dimension_checks.append({
                "dimension": dimension_name,
                "indicator": display_metric,
                "source_aggregate": source_value,
                "dashboard_sum": json_number(display_total),
                "match": Decimal(str(source_value)) == display_total,
            })
    validation["dashboard_reconciliation"] = {
        "period": latest,
        "excluded": dashboard["meta"]["excludedIndicators"],
        "kpi_checks": kpi_checks,
        "dimension_checks": dimension_checks,
        "all_match": all(item["match"] for item in [*kpi_checks, *dimension_checks]),
    }
    return dashboard, history, validation


def totals_from_row(row: dict[str, Any]) -> Totals:
    totals = Totals()
    for metric in METRICS:
        totals.values[metric] = Decimal(str(row.get(metric, 0)))
    return totals


def seed_from_history(
    aggregator: Aggregator,
    history: dict[str, Any],
    validation: dict[str, Any],
) -> None:
    for row in history.get("monthly", []):
        aggregator.monthly[row["period"]] = totals_from_row(row)
    for dimension_name, rows in history.get("dimensions", {}).items():
        if dimension_name not in aggregator.dimension:
            continue
        for row in rows:
            aggregator.dimension[dimension_name][(row["period"], row["label"])] = totals_from_row(row)
    for scope, periods in validation.get("regional_scope", {}).items():
        if scope not in aggregator.scope_monthly:
            continue
        for period, row in periods.items():
            aggregator.scope_monthly[scope][period] = totals_from_row(row)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archive", type=Path)
    parser.add_argument("--config", type=Path, default=Path("config/data-standardization.json"))
    parser.add_argument("--output-dir", type=Path, default=Path("data/generated"))
    parser.add_argument("--normalized-csv", action="store_true")
    parser.add_argument("--publish-dashboard", action="store_true")
    parser.add_argument("--base-history", type=Path)
    parser.add_argument("--base-validation", type=Path)
    args = parser.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8"))
    before_hash = sha256(args.archive)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_handle = None
    csv_writer = None
    if args.normalized_csv:
        csv_handle = (args.output_dir / "normalized-records.csv").open("w", encoding="utf-8-sig", newline="")
        csv_writer = csv.DictWriter(csv_handle, fieldnames=[
            "dataset_type", "period", "province", "district", "industry_major",
            "industry_middle", "industry_detail", "age_group", "business_age", *METRICS,
            "source_file", "source_row",
        ])
        csv_writer.writeheader()

    aggregator = Aggregator(
        config.get("age_group_mapping"),
        config["target_province"],
        config["target_districts"],
    )
    file_reports: list[dict[str, Any]] = []
    pipeline_warnings: list[str] = []
    base_source_hash = ""
    if bool(args.base_history) != bool(args.base_validation):
        raise ValueError("증분 갱신 시 --base-history와 --base-validation을 함께 지정해야 합니다.")
    if args.base_history and args.base_validation:
        base_history = json.loads(args.base_history.read_text(encoding="utf-8"))
        base_validation = json.loads(args.base_validation.read_text(encoding="utf-8"))
        seed_from_history(aggregator, base_history, base_validation)
        file_reports.extend(base_validation.get("files", []))
        pipeline_warnings.extend(base_validation.get("warnings", []))
        base_source_hash = str(base_history.get("meta", {}).get("sourceSha256", ""))
    with tempfile.TemporaryDirectory(prefix="busansinbo-transform-", ignore_cleanup_errors=True) as temp_dir:
        with zipfile.ZipFile(args.archive, "r") as archive:
            archive.extractall(temp_dir)
        root = Path(temp_dir)
        balance_files, accident_files, select_warnings = choose_source_files(root)
        pipeline_warnings.extend(select_warnings)
        incoming_periods = {period for _, period in balance_files}
        incoming_periods.update(period for _, period in accident_files if period)
        duplicate_periods = sorted(incoming_periods.intersection(aggregator.monthly))
        if duplicate_periods:
            raise ValueError(f"기존 이력과 중복되는 기준월: {duplicate_periods}")
        sources = []
        for accident_file, fixed_period in accident_files:
            required = list(config["required_transaction_headers"])
            if fixed_period:
                required = [name for name in required if name != "기준년월"]
            sources.append((accident_file, "transaction", fixed_period, required))
        sources += [(path, "balance", period, config["required_balance_headers"]) for path, period in balance_files]

        for file_index, (path, dataset_type, fixed_period, required) in enumerate(sources, start=1):
            source_totals = Totals()
            converted_totals = Totals()
            total_rows = selected_rows = converted_rows = 0
            regional_rows = {"busan_16_districts": 0, "outside_busan": 0, "busan_other": 0}
            for source_row, record in records_from_workbook(
                path, dataset_type, required, config["source_headers"], fixed_period
            ):
                total_rows += 1
                selected_rows += 1
                if record.get("province") != config["target_province"]:
                    regional_rows["outside_busan"] += 1
                elif record.get("district") in config["target_districts"]:
                    regional_rows["busan_16_districts"] += 1
                else:
                    regional_rows["busan_other"] += 1
                source_totals.add(record)
                aggregator.add(record)
                converted_totals.add(record)
                converted_rows += 1
                if csv_writer:
                    csv_record = {
                        key: json_number(value) if isinstance(value, Decimal) else value
                        for key, value in record.items()
                    }
                    csv_record.update({"source_file": path.name, "source_row": source_row})
                    csv_writer.writerow(csv_record)
            file_reports.append({
                "file": path.relative_to(root).as_posix(),
                "dataset_type": dataset_type,
                "period": fixed_period,
                "source_data_rows": total_rows,
                "selected_rows": selected_rows,
                "converted_rows": converted_rows,
                "regional_rows": regional_rows,
                "source_totals": source_totals.as_json(),
                "converted_totals": converted_totals.as_json(),
            })
            print(
                f"[{file_index}/{len(sources)}] {path.name}: 전체 {selected_rows:,}행 "
                f"(부산16개구군 {regional_rows['busan_16_districts']:,}, "
                f"부산외 {regional_rows['outside_busan']:,}, 부산기타 {regional_rows['busan_other']:,})",
                flush=True,
            )

    if csv_handle:
        csv_handle.close()
    after_hash = sha256(args.archive)
    if before_hash != after_hash:
        raise RuntimeError("원본 ZIP 해시가 작업 전후 달라졌습니다.")

    source_hash = before_hash
    if base_source_hash:
        source_hash = hashlib.sha256(f"{base_source_hash}:{before_hash}".encode("utf-8")).hexdigest()
    dashboard, history, validation = build_outputs(
        aggregator, config, source_hash, file_reports, pipeline_warnings
    )
    validation["source_unchanged"] = before_hash == after_hash
    for filename, payload in (
        ("dashboard-data.json", dashboard),
        ("dashboard-history.json", history),
        ("validation-report.json", validation),
    ):
        (args.output_dir / filename).write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    if args.publish_dashboard:
        Path("data/dashboard-data.json").write_text(
            json.dumps(dashboard, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(f"완료: {args.output_dir}")


if __name__ == "__main__":
    main()
