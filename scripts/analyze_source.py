#!/usr/bin/env python3
"""Read-only structural profiler for the source Excel archive."""

from __future__ import annotations

import argparse
import json
import re
import tempfile
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def value_type(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def profile_sheet(ws, preview_rows: int = 30, preview_cols: int = 40) -> dict[str, Any]:
    rows: list[list[Any]] = []
    type_counts: list[Counter] = [Counter() for _ in range(min(ws.max_column, preview_cols))]
    number_formats: list[Counter] = [Counter() for _ in range(min(ws.max_column, preview_cols))]
    best_header = {"row": None, "non_empty": -1, "values": []}

    for row_idx, row in enumerate(
        ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, preview_rows),
            max_col=min(ws.max_column, preview_cols),
        ),
        start=1,
    ):
        values = [json_value(cell.value) for cell in row]
        rows.append(values)
        non_empty = sum(v not in (None, "") for v in values)
        if non_empty > best_header["non_empty"]:
            best_header = {"row": row_idx, "non_empty": non_empty, "values": values}
        for idx, cell in enumerate(row):
            type_counts[idx][value_type(cell.value)] += 1
            if cell.value is not None and cell.number_format:
                number_formats[idx][cell.number_format] += 1

    return {
        "title": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "merged_ranges": [],
        "header_candidate": best_header,
        "preview": rows,
        "column_profiles": [
            {
                "column_index": i + 1,
                "types": dict(type_counts[i]),
                "number_formats": dict(number_formats[i].most_common(5)),
            }
            for i in range(len(type_counts))
        ],
    }


def period_from_name(name: str) -> str | None:
    match = re.search(r"_(\d{2})(\d{2})\.xlsx$", name)
    if not match:
        return None
    year, month = int(match.group(1)), int(match.group(2))
    return f"20{year:02d}-{month:02d}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archive", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--preview-rows", type=int, default=30)
    parser.add_argument("--preview-cols", type=int, default=40)
    args = parser.parse_args()

    source_stat = args.archive.stat()
    report: dict[str, Any] = {
        "source_archive": str(args.archive.resolve()),
        "source_size": source_stat.st_size,
        "source_modified": datetime.fromtimestamp(source_stat.st_mtime).isoformat(),
        "files": [],
        "errors": [],
    }

    with tempfile.TemporaryDirectory(prefix="busansinbo-source-", ignore_cleanup_errors=True) as temp_dir:
        with zipfile.ZipFile(args.archive, "r") as archive:
            archive.extractall(temp_dir)
        files = sorted(Path(temp_dir).rglob("*.xlsx"))

        for index, path in enumerate(files, start=1):
            relative = path.relative_to(temp_dir).as_posix()
            item: dict[str, Any] = {
                "path": relative,
                "category": path.parent.name,
                "period": period_from_name(path.name),
                "size": path.stat().st_size,
                "sheets": [],
            }
            try:
                workbook = load_workbook(path, read_only=True, data_only=False)
                item["sheet_names"] = workbook.sheetnames
                for sheet in workbook.worksheets:
                    item["sheets"].append(
                        profile_sheet(sheet, args.preview_rows, args.preview_cols)
                    )
                workbook.close()
            except Exception as exc:  # preserve the rest of the inventory
                item["error"] = f"{type(exc).__name__}: {exc}"
                report["errors"].append({"path": relative, "error": item["error"]})
            report["files"].append(item)
            print(f"[{index}/{len(files)}] {relative}", flush=True)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
